import tkinter as tk
from tkinter import messagebox
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import date
import os
import sys

# ─────────────────────────────────────────────
# CONFIGURATION — Edit these values as needed
# ─────────────────────────────────────────────
EXCEL_PATH   = r"C:\Users\heman\OneDrive\Desktop\files\Daily_Work_Log.xlsx"  # <-- Change this
DURATION     = "Full Day"
TASK         = "Auditing"
SUBTASKS     = "Financial Statement Audit"
ASSOCIATE    = "Raj Kumar"
WEEKEND_TASK = "Weekend"
# ─────────────────────────────────────────────

# Colors
C_DARK   = "#1B2E4B"
C_WHITE  = "#FFFFFF"
C_BG     = "#F4F6F9"
C_CARD   = "#EEF2F7"
C_MUTED  = "#8A94A6"
C_TEXT   = "#1A1A2E"
C_BORDER = "#D8DEE9"
C_RED    = "#E24B4A"
C_GREEN  = "#1D9E75"

HEADERS = ["Sr No", "Date", "Duration", "Task", "Subtasks", "Other Task", "Associate Name"]
COL_WIDTHS = [8, 14, 12, 20, 28, 22, 18]


def get_sheet_name(for_date=None):
    """Returns sheet name like Apr26, May26 for the given date."""
    d = for_date or date.today()
    return d.strftime("%b%y")  # e.g. Jun26


def create_sheet_headers(ws):
    """Add styled header row to a new sheet."""
    header_font  = Font(bold=True, color="FFFFFF", name="Arial", size=11)
    header_fill  = PatternFill("solid", start_color="1B2E4B")
    center       = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin         = Side(style="thin", color="AAAAAA")
    border       = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col, (h, w) in enumerate(zip(HEADERS, COL_WIDTHS), 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font      = header_font
        cell.fill      = header_fill
        cell.alignment = center
        cell.border    = border
        ws.column_dimensions[get_column_letter(col)].width = w

    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"


def get_or_create_sheet(wb, sheet_name):
    """Return existing sheet or create a new one with headers."""
    if sheet_name in wb.sheetnames:
        return wb[sheet_name]
    ws = wb.create_sheet(title=sheet_name)
    create_sheet_headers(ws)
    return ws


def get_next_sr_no(ws):
    """Find next serial number in this sheet."""
    for row in range(ws.max_row, 1, -1):
        val = ws.cell(row=row, column=1).value
        if val is not None and str(val).strip() != "":
            try:
                return int(val) + 1
            except ValueError:
                pass
    return 1


def already_logged_today(ws, today_str):
    """Check if today's date already exists in this sheet."""
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[1] and str(row[1]).strip() == today_str:
            return True
    return False


def style_row(ws, row_num, is_alt):
    thin   = Side(style="thin", color="DDDDDD")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    fill   = PatternFill("solid", start_color="EEF2F7" if is_alt else "FFFFFF")
    for col in range(1, 8):
        cell = ws.cell(row=row_num, column=col)
        cell.font      = Font(name="Arial", size=10)
        cell.fill      = fill
        cell.border    = border
        cell.alignment = Alignment(
            horizontal="center" if col in [1, 2, 3] else "left",
            vertical="center", wrap_text=True
        )
    ws.row_dimensions[row_num].height = 22


def write_to_excel(other_task_value):
    today      = date.today()
    today_str  = today.strftime("%#m/%#d/%Y") if sys.platform == "win32" else today.strftime("%-m/%-d/%Y")
    is_weekend = today.weekday() >= 5
    sheet_name = get_sheet_name(today)

    if not os.path.exists(EXCEL_PATH):
        messagebox.showerror("File Not Found",
            f"Excel file not found:\n{EXCEL_PATH}\n\nPlease update EXCEL_PATH in the script.")
        return False

    wb = openpyxl.load_workbook(EXCEL_PATH)
    ws = get_or_create_sheet(wb, sheet_name)

    if already_logged_today(ws, today_str):
        messagebox.showinfo("Already Logged", f"Work for {today_str} is already logged in sheet '{sheet_name}'!")
        wb.close()
        return True

    sr_no    = get_next_sr_no(ws)
    new_row  = ws.max_row + 1

    if is_weekend:
        row_data = [sr_no, today_str, DURATION, WEEKEND_TASK, "", "", ""]
    else:
        row_data = [sr_no, today_str, DURATION, TASK, SUBTASKS, other_task_value.strip(), ASSOCIATE]

    for col, val in enumerate(row_data, 1):
        ws.cell(row=new_row, column=col, value=val)

    style_row(ws, new_row, new_row % 2 == 0)
    wb.save(EXCEL_PATH)
    return True


def get_entry_count():
    """Get next Sr No from today's monthly sheet."""
    sheet_name = get_sheet_name()
    if not os.path.exists(EXCEL_PATH):
        return "—"
    try:
        wb = openpyxl.load_workbook(EXCEL_PATH)
        ws = get_or_create_sheet(wb, sheet_name)
        sr = get_next_sr_no(ws)
        wb.close()
        return sr
    except Exception:
        return "—"


def show_popup():
    today      = date.today()
    day_name   = today.strftime("%A")
    date_str   = today.strftime("%B %d, %Y")
    today_fmt  = today.strftime("%#m/%#d/%Y") if sys.platform == "win32" else today.strftime("%-m/%-d/%Y")
    is_weekend = today.weekday() >= 5
    sheet_name = get_sheet_name(today)

    # Check if already logged before showing window
    if os.path.exists(EXCEL_PATH):
        try:
            wb = openpyxl.load_workbook(EXCEL_PATH)
            if sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                if already_logged_today(ws, today_fmt):
                    wb.close()
                    messagebox.showinfo("Already Logged", f"Work for {date_str} is already logged!")
                    return
            wb.close()
        except Exception:
            pass

    root = tk.Tk()
    root.title("Daily Work Logger")
    root.resizable(False, False)
    root.configure(bg=C_BG)

    W, H = 420, 520 if not is_weekend else 400
    sw, sh = root.winfo_screenwidth(), root.winfo_screenheight()
    root.geometry(f"{W}x{H}+{(sw - W) // 2}+{(sh - H) // 2}")
    root.lift()
    root.attributes("-topmost", True)

    # ── HEADER ──────────────────────────────────────
    header = tk.Frame(root, bg=C_DARK, height=80)
    header.pack(fill="x")
    header.pack_propagate(False)

    icon_box = tk.Frame(header, bg="#2A3F5F", width=44, height=44)
    icon_box.place(x=18, y=18)
    tk.Label(icon_box, text="📋", font=("Segoe UI Emoji", 18),
             bg="#2A3F5F").place(relx=0.5, rely=0.5, anchor="center")

    tk.Label(header, text="Daily Work Logger",
             font=("Segoe UI", 13, "bold"), bg=C_DARK, fg=C_WHITE).place(x=74, y=18)
    tk.Label(header, text=f"{day_name}, {date_str}",
             font=("Segoe UI", 10), bg=C_DARK, fg="#8BA3C7").place(x=74, y=42)

    # Sheet badge top right
    badge_frame = tk.Frame(header, bg="#2A3F5F", padx=8, pady=3)
    badge_frame.place(x=W - 90, y=26)
    tk.Label(badge_frame, text=f"📄 {sheet_name}",
             font=("Segoe UI", 9), bg="#2A3F5F", fg="#8BA3C7").pack()

    # ── BODY ────────────────────────────────────────
    body = tk.Frame(root, bg=C_BG)
    body.pack(fill="both", expand=True, padx=18, pady=16)

    if is_weekend:
        wk_card = tk.Frame(body, bg=C_CARD, pady=20)
        wk_card.pack(fill="x", pady=(10, 0))
        tk.Label(wk_card, text="🎉", font=("Segoe UI Emoji", 30), bg=C_CARD).pack()
        tk.Label(wk_card, text="It's the weekend!",
                 font=("Segoe UI", 13, "bold"), bg=C_CARD, fg=C_TEXT).pack(pady=(6, 2))
        tk.Label(wk_card, text="Logging as Weekend — no input needed.",
                 font=("Segoe UI", 10), bg=C_CARD, fg=C_MUTED).pack()
        tk.Label(wk_card, text=f"Sheet: {sheet_name}",
                 font=("Segoe UI", 9), bg=C_CARD, fg=C_MUTED).pack(pady=(4, 0))

        def submit_weekend():
            if write_to_excel(""):
                show_success(root, "Weekend logged! Enjoy your break 🎉")

        tk.Button(body, text="  ✓  Log Weekend & Close",
                  command=submit_weekend,
                  bg=C_DARK, fg=C_WHITE,
                  font=("Segoe UI", 11, "bold"),
                  relief="flat", cursor="hand2",
                  activebackground="#253D60", activeforeground=C_WHITE,
                  padx=10, pady=10).pack(fill="x", pady=(20, 0))

    else:
        # ── Sheet info strip ──
        sheet_strip = tk.Frame(body, bg="#E8EDF5", pady=6, padx=12)
        sheet_strip.pack(fill="x", pady=(0, 10))
        tk.Label(sheet_strip, text=f"📄  Logging to sheet:  ",
                 font=("Segoe UI", 9), bg="#E8EDF5", fg=C_MUTED).pack(side="left")
        tk.Label(sheet_strip, text=sheet_name,
                 font=("Segoe UI", 9, "bold"), bg="#E8EDF5", fg=C_DARK).pack(side="left")
        sr = get_entry_count()
        tk.Label(sheet_strip, text=f"   •   Entry #{sr}",
                 font=("Segoe UI", 9), bg="#E8EDF5", fg=C_MUTED).pack(side="left")

        # ── Auto-fill card ──
        af_frame = tk.Frame(body, bg=C_CARD, pady=10, padx=14)
        af_frame.pack(fill="x")

        lock_row = tk.Frame(af_frame, bg=C_CARD)
        lock_row.pack(fill="x", anchor="w")
        tk.Label(lock_row, text="🔒  AUTO-FILLED FIELDS",
                 font=("Segoe UI", 8, "bold"), bg=C_CARD, fg=C_MUTED).pack(side="left")

        tk.Frame(af_frame, bg=C_BORDER, height=1).pack(fill="x", pady=(6, 8))

        grid = tk.Frame(af_frame, bg=C_CARD)
        grid.pack(fill="x")

        fields = [("Duration", DURATION), ("Task", TASK),
                  ("Subtasks", SUBTASKS), ("Associate", ASSOCIATE)]

        for i, (label, val) in enumerate(fields):
            cell = tk.Frame(grid, bg=C_CARD)
            cell.grid(row=i // 2, column=i % 2, sticky="w", padx=(0, 20), pady=3)
            tk.Label(cell, text=label,
                     font=("Segoe UI", 9), bg=C_CARD, fg=C_MUTED).pack(anchor="w")
            tk.Label(cell, text=val,
                     font=("Segoe UI", 10, "bold"), bg=C_CARD, fg=C_TEXT).pack(anchor="w")

        # ── Other Task input ──
        tk.Frame(body, bg=C_BORDER, height=1).pack(fill="x", pady=14)

        tk.Label(body, text="✏️  What else did you work on today?",
                 font=("Segoe UI", 10, "bold"), bg=C_BG, fg=C_TEXT).pack(anchor="w")

        entry_var   = tk.StringVar()
        entry_frame = tk.Frame(body, bg=C_BORDER, pady=1, padx=1)
        entry_frame.pack(fill="x", pady=(8, 0))

        entry = tk.Entry(entry_frame, textvariable=entry_var,
                         font=("Segoe UI", 11), relief="flat",
                         bg=C_WHITE, fg=C_TEXT, insertbackground=C_DARK)
        entry.pack(fill="x", ipady=9, padx=1)
        entry.focus_set()

        tk.Label(body, text="e.g.  Excel,  Report,  Client meeting...",
                 font=("Segoe UI", 9), bg=C_BG, fg=C_MUTED).pack(anchor="w", pady=(4, 0))

        error_lbl = tk.Label(body, text="", font=("Segoe UI", 9), bg=C_BG, fg=C_RED)
        error_lbl.pack(anchor="w")

        def on_entry_change(*_):
            entry_frame.config(bg=C_DARK)
            error_lbl.config(text="")

        entry_var.trace_add("write", on_entry_change)

        # ── Buttons ──
        tk.Frame(body, bg=C_BORDER, height=1).pack(fill="x", pady=12)

        btn_row = tk.Frame(body, bg=C_BG)
        btn_row.pack(fill="x")

        def submit(event=None):
            val = entry_var.get().strip()
            if not val:
                entry_frame.config(bg=C_RED)
                error_lbl.config(text="⚠  Please enter your Other Task for today.")
                entry.focus_set()
                return
            if write_to_excel(val):
                show_success(root, f"Logged to {sheet_name} ✓")

        def skip():
            if write_to_excel(""):
                show_success(root, f"Logged to {sheet_name} (no Other Task) ✓")

        root.bind("<Return>", submit)

        tk.Button(btn_row, text="  💾  Save & Log",
                  command=submit,
                  bg=C_DARK, fg=C_WHITE,
                  font=("Segoe UI", 11, "bold"),
                  relief="flat", cursor="hand2",
                  activebackground="#253D60", activeforeground=C_WHITE,
                  padx=12, pady=9).pack(side="left", fill="x", expand=True, padx=(0, 6))

        tk.Button(btn_row, text="Skip",
                  command=skip,
                  bg=C_BG, fg=C_MUTED,
                  font=("Segoe UI", 10),
                  relief="flat", cursor="hand2",
                  highlightthickness=1,
                  highlightbackground=C_BORDER,
                  activebackground=C_CARD,
                  padx=14, pady=9).pack(side="left")

    root.mainloop()


def show_success(parent, message):
    top = tk.Toplevel(parent)
    top.overrideredirect(True)
    top.configure(bg=C_GREEN)
    pw, py, pw2 = parent.winfo_x(), parent.winfo_y(), parent.winfo_width()
    top.geometry(f"300x70+{pw + (pw2 - 300) // 2}+{py + 20}")
    top.attributes("-topmost", True)
    tk.Label(top, text=message,
             font=("Segoe UI", 11, "bold"),
             bg=C_GREEN, fg=C_WHITE, pady=22).pack(fill="both", expand=True)
    top.after(1800, lambda: (top.destroy(), parent.destroy()))


if __name__ == "__main__":
    show_popup()